"""生成脱敏版演示排课表（schedule.demo.xlsx）。

背景：仓库根目录的 schedule.xlsx 是真实教学数据，被 .gitignore 排除（"kept local"）。
若要把它随项目发布，必须先脱敏，否则会公开：
  - 「教师联系表」中的真实手机号与邮箱（PII，实测 6 个手机号）
  - 「学生所在系系办电话」中的院系办公电话
  - 「任课老师安排表」中的学院名称（闽江学院）与内部排课备注
  - 25 个真实人名（经 /status、/teachers、/parallel/overview 等只读接口可被任意访客读取）

脱敏策略（保结构、去身份）：
  1. 丢弃应用**不读取**的敏感工作表：学生所在系系办电话、任课老师安排表
     （解析器只读「大学物理实验安排表 (周) 」与「教师联系表」，见 backend/app/utils/excel_parser.py）
  2. 全文姓名替换：真实姓名 → 稳定化名「教师A、教师B…」。
     对**所有工作表的所有单元格**做替换，覆盖「教师」「实验筹备人员」等列，
     以及 "钟蔚\\n林珠云"（一格多名）、"鄢仁文吴琼"（无分隔粘连）、"郑辉（落任务排前面）"（带括注）。
     长名优先替换，避免短名截断长名。
  3. 「教师联系表」清空 电话 / 邮箱 两列（保留 姓名 + 备注 D 列，因为 get_teacher_roles 会读 D 列）
  4. 清除文档属性（creator / lastModifiedBy / title），避免元数据里残留姓名
  5. 班级名、教室、实验项目、周次时段等教学结构**原样保留**，保证演示效果等价

用法：
    python scripts/make_demo_schedule.py                 # schedule.xlsx → schedule.demo.xlsx
    python scripts/make_demo_schedule.py <in> <out>
"""

import re
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRC = REPO_ROOT / "schedule.xlsx"
DEFAULT_DST = REPO_ROOT / "schedule.demo.xlsx"

# 解析器实际读取的两张表
SHEET_WEEK = "大学物理实验安排表 (周) "   # 注意结尾有空格
SHEET_CLASS = "大学物理实验安排表 (班级) "
SHEET_CONTACT = "教师联系表"

# 应用不读取、且含敏感信息 → 整表丢弃
DROP_SHEETS = ["学生所在系系办电话", "任课老师安排表"]

# 收集"真实姓名"的来源列（0-based）。收集范围可宽于替换范围：
# 替换是全表全局进行的，收集只决定"哪些字符串算人名"。
NAME_SOURCE_COLUMNS = {
    SHEET_WEEK: [7, 9],       # H 教师、J 实验筹备人员
    SHEET_CLASS: [7, 9],      # H 教师、J 实验筹备人员
    SHEET_CONTACT: [0],       # A 姓名
}

# 表头/字段名黑名单：这些 2-4 字中文串不是人名，必须排除，
# 否则会产生「教师 → 教师I」这类自指化名，替换后仍是残留。
HEADER_BLOCKLIST = {
    "姓名", "教师", "电话", "邮箱", "备注", "实验筹备人员", "任课老师",
    "班级", "场地", "项目", "组别", "周次", "周别", "时段", "课程类别",
    "人数", "合班", "合班人数", "分组数", "起始周", "轮数", "学生学院",
    "公共课程名称", "周学时", "合班情况", "实验课时间", "绪论课时间",
    "教学秘书联系电话", "班级序号", "周序号", "序号", "实验筹备人",
}

# 「教师联系表」中需要清空的敏感列（0-based）：B 电话、C 邮箱
CONTACT_SCRUB_COLUMNS = [1, 2]

# 11 位大陆手机号
PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
# 姓名后常见的括注，如「郑辉（落任务排前面）」→ 取出「郑辉」
PAREN_NOTE_RE = re.compile(r"[（(].*?[)）]")
# 纯中文姓名
CN_NAME_RE = re.compile(r"^[\u4e00-\u9fa5]{2,4}$")


def _clean_name(raw: str) -> str:
    """去掉括注与首尾空白，得到纯姓名。"""
    return PAREN_NOTE_RE.sub("", str(raw)).strip()


def _is_name_like(candidate: str) -> bool:
    if not candidate or candidate in HEADER_BLOCKLIST:
        return False
    return bool(CN_NAME_RE.fullmatch(candidate))


def collect_real_names(wb: openpyxl.Workbook) -> set:
    """从已知来源列收集真实姓名（排除表头词）。"""
    names = set()
    for sheet_name, cols in NAME_SOURCE_COLUMNS.items():
        if sheet_name not in wb.sheetnames:
            continue
        for row in wb[sheet_name].iter_rows(values_only=True):
            for c in cols:
                if c >= len(row) or not row[c]:
                    continue
                for piece in str(row[c]).replace("\r", "\n").split("\n"):
                    n = _clean_name(piece)
                    if _is_name_like(n):
                        names.add(n)
    return names


def build_alias_map(names: set) -> dict:
    """姓名 → 化名。按姓名排序保证跨次运行结果稳定。"""
    return {n: f"教师{chr(ord('A') + i)}" for i, n in enumerate(sorted(names))}


def anonymize_text(text: str, alias: dict) -> str:
    """对单格文本做姓名替换（长名优先，避免短名截断长名）。"""
    if text is None or text == "":
        return text
    out = str(text)
    for real in sorted(alias, key=len, reverse=True):
        if real in out:
            out = out.replace(real, alias[real])
    return out


def main() -> int:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_DST

    if not src.is_file():
        print(f"[ERR] 源文件不存在：{src}")
        return 1

    wb = openpyxl.load_workbook(src)

    names = collect_real_names(wb)
    alias = build_alias_map(names)
    print(f"收集到真实姓名 {len(names)} 个 → 已生成化名映射")
    print("  明细: " + "、".join(f"{k}→{v}" for k, v in sorted(alias.items())))

    # 1) 丢弃应用不读取的敏感工作表
    for sh in DROP_SHEETS:
        if sh in wb.sheetnames:
            del wb[sh]
            print(f"已删除工作表：{sh}")

    # 2) 全表全局替换姓名（覆盖所有工作表、所有单元格）
    replaced_cells = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    new = anonymize_text(cell.value, alias)
                    if new != cell.value:
                        cell.value = new
                        replaced_cells += 1
    print(f"已替换 {replaced_cells} 个单元格中的姓名")

    # 3) 「教师联系表」清空电话/邮箱
    if SHEET_CONTACT in wb.sheetnames:
        scrubbed = 0
        for row in wb[SHEET_CONTACT].iter_rows(min_row=1):
            for c in CONTACT_SCRUB_COLUMNS:
                if c < len(row) and row[c].value:
                    row[c].value = None
                    scrubbed += 1
        print(f"已清空「{SHEET_CONTACT}」{scrubbed} 个电话/邮箱单元格")

    # 4) 清除文档属性中的姓名等元数据
    props = wb.properties
    props.creator = None
    props.lastModifiedBy = None
    props.title = None
    props.subject = None
    props.description = None

    wb.save(dst)
    print(f"已生成：{dst}")

    # --- 自检：确认残留 PII = 0 ---
    check = openpyxl.load_workbook(dst, data_only=True)
    leftover_names = set()
    leftover_phones = []
    for ws in check.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if not v:
                    continue
                s = str(v)
                for p in PHONE_RE.findall(s):
                    leftover_phones.append((ws.title, p))
                for n in names:
                    if n in s:
                        leftover_names.add((ws.title, n))
    print("--- 自检 ---")
    print(f"剩余工作表：{check.sheetnames}")
    print(f"残留真实姓名：{len(leftover_names)} {sorted(leftover_names) if leftover_names else ''}")
    print(f"残留手机号  ：{len(leftover_phones)} {leftover_phones if leftover_phones else ''}")
    ok = not leftover_names and not leftover_phones
    print("结论：" + ("PASS 脱敏干净" if ok else "FAIL 仍有残留，不可发布"))
    return 0 if ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
