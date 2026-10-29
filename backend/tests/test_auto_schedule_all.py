"""全校一键排课（R-59）：全表覆盖 + 分片不互相撞车 + 初稿可导出

缺口：`/auto-schedule` 长时间 `max_sections=30`，真实表 171 个教学任务里
**141 个被静默略过**（文案写了"取前 30"，但教秘要的是全校排完）。
单次 171 个任务实测 6 秒时限内连可行解都拿不到（UNKNOWN），所以改分片，
片与片之间必须回填已排定的 教室/教师/班级 占用 —— 这是本批最容易写错的地方：
少回填一项，两片就会把同一个格子排给两个人，而单看每片都"成功"。
"""

import asyncio
import io
import os
import sys
from collections import Counter

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.engine.auto_scheduler import AutoScheduler
from app.master_data import MasterData
from app.utils.excel_parser import ScheduleDatabase, ScheduleEntry

import app.api.reschedule as api

REAL_XLSX = os.path.join(os.path.dirname(__file__), "..", "..", "schedule.xlsx")


def _db(entries):
    d = ScheduleDatabase()
    d.entries = list(entries)
    d.skipped = []
    d.loaded = True
    return d


def _entry(week, weekday, slot, cls, room, teacher, experiment):
    return ScheduleEntry(week, weekday, slot, cls, "1组", experiment, room, teacher)


def _synthetic_db(n_teachers=4, n_classes=5, weeks=6):
    """无重复 (教师, 项目, 班级) 的合成表：任务数 = 教师 × 班级。"""
    entries = []
    for w in range(1, weeks + 1):
        for ti in range(n_teachers):
            for ci in range(n_classes):
                entries.append(_entry(
                    w, "周一", "567", f"班{ci}", f"A{400 + ci}",
                    f"师{ti}", f"项目{ti}"))
    return _db(entries)


# ---------------------------------------------------------------- 全表覆盖

def test_solve_all_covers_every_section():
    """合成 20 个任务 → 分 4 片（每片 5 个）全部排定，不是前 N 个。"""
    db = _synthetic_db()  # 4 教师 × 5 班 = 20 个 (教师, 项目, 班级)
    r = AutoScheduler(db).solve_all(chunk_size=5)

    assert r["success"] is True
    assert r["total_sections"] == 20
    assert r["solved"] == 20, f"只排了 {r['solved']}/20"
    assert len(r["chunks"]) == 4
    assert all(c["solved"] == c["size"] for c in r["chunks"]), r["chunks"]
    assert not r["conflicts"]


def test_chunks_do_not_double_book_across_boundaries():
    """分片最容易犯的错：第 2 片看不见第 1 片排过的格子，于是排到同一格。

    判据按 场地 / 教师 / 班级 各算一遍 —— 三个维度都必须零重复。
    """
    db = _synthetic_db()
    r = AutoScheduler(db).solve_all(chunk_size=3)

    rooms = Counter((s["week"], s["weekday"], s["time_slot"], s["room"]) for s in r["schedule"])
    teachers = Counter((s["week"], s["weekday"], s["time_slot"], s["teacher"]) for s in r["schedule"])
    classes = Counter((s["week"], s["weekday"], s["time_slot"], s["class_name"]) for s in r["schedule"])
    assert max(rooms.values()) == 1, "同一格同一场地被排了两次（片间未回填 room_busy）"
    assert max(teachers.values()) == 1, "同一格同一教师被排了两次（片间未回填 teacher_busy）"
    assert max(classes.values()) == 1, "同一格同一班级被排了两次（片间未回填 class_busy）"


def test_empty_db_reports_no_tasks():
    r = AutoScheduler(_db([])).solve_all(chunk_size=10)
    assert r["success"] is False and "没有可排课" in r["error"]


# ---------------------------------------------------------------- R-72 组展开

def _entry_g(week, weekday, slot, cls, room, teacher, experiment, group):
    return ScheduleEntry(week, weekday, slot, cls, group, experiment, room, teacher)


def test_expand_groups_same_week_contiguous_and_avoids_kept():
    """段级求解 → 组级产出：k 个组必须同周、同教室、时间连续，且不得排到
    "不参与重排"的绪论/考试格子上（教室与班级两个维度都验）。"""
    from app.engine.auto_scheduler import TIME_SLOTS, WEEKDAYS

    sched = AutoScheduler(_db([]))
    row = {"week": 1, "weekday": "周一", "time_slot": "567", "room": "A410",
           "teacher": "师0", "course": "项目0", "class_name": "班0"}
    groups_of = {("师0", "项目0", "班0"): ["组1", "组2", "组3"]}

    # 保留课占用：周一 678/91011 的 A410（教室维度）、周二 567 的班0（班级维度）
    room_busy = {(1, "周一", "678", "A410"), (1, "周一", "91011", "A410")}
    class_busy = {(1, "周二", "567", "班0")}

    out, gaps, total = sched._expand_groups(
        [row], groups_of, set(room_busy), set(), set(class_busy))

    assert total == 3
    assert not gaps and len(out) == 3
    assert [r["group"] for r in out] == ["组1", "组2", "组3"]
    assert all(r["week"] == 1 and r["room"] == "A410" for r in out)
    # 时间顺序连续
    timeline = [(d, t) for d in WEEKDAYS for t in TIME_SLOTS]
    pos = {c: i for i, c in enumerate(timeline)}
    seq = sorted(pos[(r["weekday"], r["time_slot"])] for r in out)
    assert seq == list(range(seq[0], seq[0] + 3)), seq
    # 不撞保留课
    for r in out:
        cell = (r["week"], r["weekday"], r["time_slot"])
        assert cell + (r["room"],) not in room_busy
        assert cell + (r["class_name"],) not in class_busy


def test_expand_groups_falls_back_to_other_week_when_current_full():
    """本周凑不出 k 个连续格时整段换周，而不是硬记缺口（R-72 第 4 步）。

    构造：第 1 周 A410 除段自己那一格外全部占满（k=3 排不进），第 2 周全空。
    期望：3 组整体落到第 2 周，且连续、无缺口。
    """
    from app.engine.auto_scheduler import TIME_SLOTS, WEEKDAYS

    db = _db([_entry(1, "周一", "567", "班X", "A410", "师X", "项目X"),
              _entry(2, "周一", "567", "班X", "A410", "师X", "项目X")])
    sched = AutoScheduler(db)
    assert sched._get_weeks() == [1, 2]

    row = {"week": 1, "weekday": "周一", "time_slot": "567", "room": "A410",
           "teacher": "师0", "course": "项目0", "class_name": "班0"}
    groups_of = {("师0", "项目0", "班0"): ["组1", "组2", "组3"]}
    # 第 1 周除 (周一,567) 外全占，只留 1 格 → k=3 在第 1 周必失败
    room_busy = {(1, d, t, "A410") for d in WEEKDAYS for t in TIME_SLOTS
                 if (d, t) != ("周一", "567")}

    out, gaps, _total = sched._expand_groups(
        [row], groups_of, set(room_busy), set(), set())

    assert not gaps and len(out) == 3
    assert all(r["week"] == 2 for r in out), [r["week"] for r in out]
    timeline = [(d, t) for d in WEEKDAYS for t in TIME_SLOTS]
    pos = {c: i for i, c in enumerate(timeline)}
    seq = sorted(pos[(r["weekday"], r["time_slot"])] for r in out)
    assert seq == list(range(seq[0], seq[0] + 3)), seq


# ---------------------------------------------------------------- API 覆盖口径

@pytest.fixture
def api_db(monkeypatch):
    db = _synthetic_db()
    monkeypatch.setattr(api, "_db", db)
    monkeypatch.setattr(api, "_master_tasks", [])
    md = MasterData()
    monkeypatch.setattr(api, "_master_data", md)
    return db


def test_api_default_is_full_coverage(api_db):
    out = asyncio.run(api.auto_schedule(None))
    assert out["success"] is True
    d = out["data"]
    assert d["truncated"] is False
    # R-72：coverage 是「组课时（段）」双口径交付串，不再是旧的单一 `solved/total_all`。
    # 测试镜像服务端产 coverage 的同一份格式（字段名取自响应，不写死数字）。
    assert d["coverage"] == (
        f"{d['solved_groups']}/{d['total_groups']} 组课时（{d['solved']}/{d['total_sections_all']} 段）")
    assert d["solved"] == d["total_sections_all"] == 20
    assert len(d["schedule"]) == 20


def test_api_max_sections_still_truncates(api_db):
    """调试/演示仍可只要前 N 个，但必须显式标 truncated。"""
    out = asyncio.run(api.auto_schedule(api.AutoScheduleRequest(max_sections=3)))
    d = out["data"]
    assert d["truncated"] is True
    assert d["total_sections"] == 3
    assert d["total_sections_all"] == 20


# ---------------------------------------------------------------- 初稿导出

def _export_rows():
    return [
        {"week": 2, "weekday": "周一", "time_slot": "567", "room": "A410",
         "teacher": "师0", "course": "项目0", "class_name": "班0"},
        {"week": 2, "weekday": "周一", "time_slot": "678", "room": "A410",
         "teacher": "师0", "course": "项目0", "class_name": "班1"},
        {"week": 3, "weekday": "周二", "time_slot": "91011", "room": "A411",
         "teacher": "师1", "course": "项目1", "class_name": "班0"},
    ]


def _drain(resp):
    async def _collect():
        chunks = []
        async for chunk in resp.body_iterator:
            chunks.append(chunk)
        return b"".join(chunks)
    return asyncio.run(_collect())


def test_export_workbook_has_master_sheet_and_one_sheet_per_class():
    import openpyxl

    resp = asyncio.run(api.auto_schedule_export(api.AutoScheduleExportRequest(schedule=_export_rows())))
    wb = openpyxl.load_workbook(io.BytesIO(_drain(resp)))
    assert wb.sheetnames == ["总表", "班0", "班1"]
    master = [[c for c in row] for row in wb["总表"].iter_rows(values_only=True)]
    assert len(master) == 1 + 3  # 表头 + 3 条
    assert master[0] == ["周次", "星期", "时段", "场地", "教师", "实验项目", "班级"]
    assert len(list(wb["班0"].iter_rows(values_only=True))) == 1 + 2


def test_export_rejects_empty_schedule():
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as ei:
        asyncio.run(api.auto_schedule_export(api.AutoScheduleExportRequest(schedule=[])))
    assert ei.value.status_code == 400


def test_export_sanitizes_illegal_sheet_names():
    """班级名带 Excel 非法字符（真实表就有「2024级应用化学1、2班」这类）不能崩。"""
    import openpyxl

    rows = [dict(_export_rows()[0], class_name="24级/实验[1]班")]
    resp = asyncio.run(api.auto_schedule_export(api.AutoScheduleExportRequest(schedule=rows)))
    wb = openpyxl.load_workbook(io.BytesIO(_drain(resp)))
    assert "总表" in wb.sheetnames and len(wb.sheetnames) == 2
    assert "/" not in wb.sheetnames[1] and "[" not in wb.sheetnames[1]


# ---------------------------------------------------------------- 真实表规模

@pytest.mark.skipif(not os.path.exists(REAL_XLSX), reason="真实排课表不在仓库里")
def test_real_table_full_coverage():
    """真实表 171 个任务全表分片：全部排定、零跨片撞车（约 36 秒）。

    耗时说明：每片一阶段 7 秒 + 摊平 5 秒的预算，硬片会用满 —— 曾把一阶段压到
    5 秒（比"60 个一片 7 秒可解"的实测点还紧），真实表第 2 片在边界上时灵时不灵
    （单独跑 5 次里 4 次 UNKNOWN），那是时限给错了不是求解器退化（见 solve_all 注释）。
    """
    db = ScheduleDatabase()
    db.load_from_excel(os.path.abspath(REAL_XLSX))
    r = AutoScheduler(db).solve_all(chunk_size=60)

    assert r["success"] is True
    assert r["solved"] == r["total_sections"] >= 150
    # R-72：段级求解、组级产出 —— 组课时必须全覆盖（171 段 = 410 组课时），
    # 且每一行都带 group（否则落不进安排表）。
    assert r["solved_groups"] == r["total_groups"], (r["solved_groups"], r["total_groups"])
    assert all(s.get("group") for s in r["schedule"])

    rooms = Counter((s["week"], s["weekday"], s["time_slot"], s["room"]) for s in r["schedule"])
    teachers = Counter((s["week"], s["weekday"], s["time_slot"], s["teacher"]) for s in r["schedule"])
    classes = Counter((s["week"], s["weekday"], s["time_slot"], s["class_name"]) for s in r["schedule"])
    assert max(rooms.values()) == 1 and max(teachers.values()) == 1 and max(classes.values()) == 1

    # R-72：不得与"不参与重排"的绪论/考试撞格（教室/教师/班级三维）——
    # 这是 R-72 修过的真实缺陷（实测曾撞 5 处），此处锁死防回归。
    kr, kt, kc = AutoScheduler(db)._kept_busy()
    for s in r["schedule"]:
        cell = (s["week"], s["weekday"], s["time_slot"])
        assert cell + (s["room"],) not in kr, f"撞绪论教室 {s}"
        assert cell + (s["teacher"],) not in kt, f"撞绪论教师 {s}"
        assert cell + (s["class_name"],) not in kc, f"撞绪论班级 {s}"

    # R-72 闭环 #6：一键自动排课初稿 → 落地。链 #6 曾因产出缺 group 而断开，
    # 现在产出带 group，应当能过 /parallel/land 的校验并写回安排表。
    import copy
    from app import plan_landing

    plan = plan_landing.normalize_plan(r["schedule"])  # 即 POST /parallel/land 的入口
    assert len(plan) == r["total_groups"], (len(plan), r["total_groups"])
    assert all(p["group"] for p in plan), "落地方案里还有缺 group 的行"

    db2 = copy.deepcopy(db)
    land = plan_landing.apply_plan(db2, plan)         # 写回（在副本上，不污染原表）
    assert land["replaced_with"] == r["total_groups"]
    # 1) 初稿自身三维度内部零重复
    for attr in ("room", "teacher", "class_name"):
        seen = Counter((p["week"], p["weekday"], p["time_slot"], p[attr])
                      for p in plan if p.get(attr))
        assert max(seen.values()) == 1, f"初稿内部 {attr} 维度重复占用"
    # 2) 初稿不撞"不参与重排"的绪论/考试（只查新写回的 410 行，保留行本就定义它们）
    for p in plan:
        cell = (p["week"], p["weekday"], p["time_slot"])
        assert cell + (p["room"],) not in kr, f"落地后撞绪论教室 {p}"
        assert cell + (p["teacher"],) not in kt, f"落地后撞绪论教师 {p}"
        assert cell + (p["class_name"],) not in kc, f"落地后撞绪论班级 {p}"
    # 3) 落地不引入"新"碰撞：全表碰撞数不得多于原表。原表绪论/考试自身的历史重叠
    #    （如林珠云在同一格带两个班）是源数据既有问题，初稿只替换可排单元、原样保留它们，
    #    不应为此背锅；这里锁的是"初稿不制造新冲突"。
    def _collisions(entries, attr):
        c = Counter((e.week, e.weekday, e.time_slot, getattr(e, attr))
                    for e in entries if getattr(e, attr))
        return sum(v - 1 for v in c.values() if v > 1)
    for attr in ("room", "teacher", "class_name"):
        assert _collisions(db2.entries, attr) <= _collisions(db.entries, attr), \
            f"落地后 {attr} 维度新增了碰撞（原表 {_collisions(db.entries, attr)} → " \
            f"落地 {_collisions(db2.entries, attr)}）"


# ---------------------------------------------------------------- R-74 贪心/展开撞绪论
#
# 病（2026-10-28 实测复现，非推断）：CPU 饱和时 `test_real_table_full_coverage`
# 会红，红的不是覆盖率，而是「撞绪论教师」：
#     {'week': 2, 'weekday': '周一', 'time_slot': '567', 'teacher': '翁铭华', ...}
#     而 (2, '周一', '567', '翁铭华') 正是绪论课的占用格。
# 根因是**两个**缺陷叠加（缺一个都不会撞）：
#   1. `_greedy_fill` 只拿到全局累积集，而累积集从空集起步（要允许复用非绪论旧格），
#      本身**不含**绪论/考试占用 ⇒ 贪心看不见它们，把段排进了绪论格；
#   2. `_expand_groups` 收到的是 `busy | kept` 的**合并集**，而它开头会
#      `discard` 自己那一格来"释放自己"—— 合并后这一步把 kept 在那一格的记录
#      **一起抹掉**，于是展开又心安理得地把段放回绪论格。
# 为什么表现为 flaky：只有 CP-SAT 在时限内没排完（负载高 / 时间盒紧）才走贪心，
# 平时走 CP-SAT 就看不见。下面三条把两个缺陷分别锁死。

# 为什么这一条必须挂 skipif（2026-10-29 实测）：它 `load_from_excel(REAL_XLSX)`，
# 而 CI 的干净检出**不含** `schedule.xlsx`（.gitignore 排除）⇒ 直接
# `FileNotFoundError: ...\schedule.xlsx` 判红。R-74 加这三条判据时，只有
# `test_real_table_full_coverage` 带了装饰器，这条漏了 —— 于是 CI 上**一直是红的**，
# 而本地（有真实表）永远是绿的，两边看到的不是同一件事。
# 同文件的另外两条（`test_greedy_fill_never_takes_kept_cells` /
# `test_expand_release_does_not_erase_kept`）用合成表，不需要它。
@pytest.mark.skipif(not os.path.exists(REAL_XLSX), reason="真实排课表不在仓库里")
def test_greedy_path_never_collides_with_kept_courses():
    """确定性复现：把时间盒压到必然 UNKNOWN ⇒ 强制走贪心+展开路径。

    不需要靠 CPU 负载去凑 flaky —— `flatten_seconds=0.05` 下 CP-SAT 每片都拿不到解，
    路径与"机器很忙"时完全一致。
    """
    db = ScheduleDatabase()
    db.load_from_excel(os.path.abspath(REAL_XLSX))
    sched = AutoScheduler(db)
    r = sched.solve_all(chunk_size=60, flatten_seconds=0.05)

    assert r["schedule"], "贪心路径也必须产出方案（否则这条测不到东西）"
    kr, kt, kc = sched._kept_busy()
    for s in r["schedule"]:
        cell = (s["week"], s["weekday"], s["time_slot"])
        assert cell + (s["room"],) not in kr, f"贪心/展开撞绪论教室 {s}"
        assert cell + (s["teacher"],) not in kt, f"贪心/展开撞绪论教师 {s}"
        assert cell + (s["class_name"],) not in kc, f"贪心/展开撞绪论班级 {s}"
    # 产出仍是组课时（带 group），否则落不了地
    assert all(s.get("group") for s in r["schedule"])


def test_greedy_fill_never_takes_kept_cells():
    """缺陷 1 的直接单测：kept 全占满时贪心必须"排不下就留空"，不是硬塞。"""
    # 唯一教师「师0」+ 唯一教室 A410 被绪论课占满第 1 周全部 15 格
    entries = [_entry(1, d, t, "一班", "A410", "师0", "绪论")
               for d in ("周一", "周二", "周三", "周四", "周五")
               for t in ("567", "678", "91011")]
    db = _db(entries)
    sched = AutoScheduler(db)
    kr, kt, kc = sched._kept_busy()
    assert kr and kt and kc, "前提：绪论占用三维都要有"

    sec = {"teacher": "师0", "course": "声速测量", "class_name": "一班", "groups": ["1组"]}
    # 不传 kept（修复前的调用形态）：贪心把段塞进绪论格 —— 这就是线上撞法的成因
    blind = sched._greedy_fill([dict(sec)], set(), set(), set())
    assert len(blind) == 1, "前提：看不见 kept 时贪心确实会占绪论格"
    assert (1, blind[0]["weekday"], blind[0]["time_slot"], "师0") in kt

    # 传 kept：一格都不能占
    seen = sched._greedy_fill([dict(sec)], set(), set(), set(), kr, kt, kc)
    assert seen == [], f"贪心不得占用绪论格，实得 {seen}"


def test_expand_release_does_not_erase_kept():
    """缺陷 2 的直接单测：展开"释放自己那一格"不得把绪论占用一起 discard。"""
    db = _db([_entry(1, "周一", "567", "一班", "A410", "师0", "绪论")])
    sched = AutoScheduler(db)
    kr, kt, kc = sched._kept_busy()
    row = {"week": 1, "weekday": "周一", "time_slot": "567", "room": "A410",
           "teacher": "师0", "course": "声速测量", "class_name": "一班"}
    # 调用方预置本批占用（这里恰好=绪论格，模拟缺陷 1 的产物）
    busy_r, busy_t, busy_c = set(kr), set(kt), set(kc)
    out, _gaps, _total = sched._expand_groups(
        [dict(row)], {("师0", "声速测量", "一班"): ["1组"]},
        busy_r, busy_t, busy_c, kr, kt, kc)

    assert out, "换到别的格子应当能排上"
    for r in out:
        cell = (r["week"], r["weekday"], r["time_slot"])
        assert cell + (r["room"],) not in kr, f"展开撞绪论教室 {r}"
        assert cell + (r["teacher"],) not in kt, f"展开撞绪论教师 {r}"
        assert cell + (r["class_name"],) not in kc, f"展开撞绪论班级 {r}"
