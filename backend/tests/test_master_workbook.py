"""主数据台账：导出 → 编辑 → 导回（P2 最后一块闭环）

要修的洞
--------
`/export/master-data` 的 docstring 承诺"可以补列、再导回"，但它产出的是**统计台账**
（「班级数」「课次(格)」这类聚合列），而 `/upload-master` 认的是**清单型**任务表 ——
两种格式不对齐，导回只会静默产出垃圾。本文件锁的是新闭环
（`/export/master-workbook` + `/master/import`）的四条防线：

  1. **往返无损**：导出再导回，五张表逐字段相等；
  2. **编辑真的生效**（反向验证 R-33）：改过的格子必须真的写回主数据；
     如果导入是 no-op，用例 2 必须红 —— 只断言"接口返回 success"会假绿；
  3. **统计列不得变成实体名**：「班级数 = 5」不许导入出一个名叫「5」的班级；
  4. **认不出就报出来**：表头认不出 → 整页跳过并说明，绝不静默导入一半（R-34）。
"""

import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl

from app import master_store
from app.master_data import MasterData
from app.master_workbook import (_split_multi, apply_tables, build_workbook,
                                parse_workbook)
from app.utils.excel_parser import ScheduleDatabase, ScheduleEntry


def _entry(week, weekday, slot, cls, group, room, teacher="吴老师", experiment="扭摆法"):
    return ScheduleEntry(week, weekday, slot, cls, group, experiment, room, teacher)


def _base_md() -> MasterData:
    md = MasterData()
    md.build_from_entries([
        _entry(1, "周四", "567", "一班", "1组", "A410", teacher="吴老师", experiment="扭摆法"),
        _entry(2, "周四", "567", "一班", "2组", "A411", teacher="林老师", experiment="扭摆法"),
        _entry(3, "周五", "678", "二班", "1组", "A410", teacher="吴老师", experiment="声速测量"),
    ])
    md.upsert_room("A410", capacity=30)
    md.upsert_class("一班", groups=["1组", "2组"], weekday="周四",
                    time_slots=["567"], week_min=1, week_max=16)
    return md


def _parse_only(path):
    """parse_workbook 返回 (tables, report)；多数用例只关心写库结果。"""
    return parse_workbook(str(path))[0]


def _dump(md, path):
    build_workbook(md).save(path)
    return str(path)


# ---------------------------------------------------------------- 1. 往返无损

def test_roundtrip_export_then_import_is_lossless(tmp_path):
    """导出 → 导入到一张空表 → 五张实体表逐字段相等。"""
    md = _base_md()
    p = _dump(md, tmp_path / "wb.xlsx")

    tables, report = parse_workbook(p)
    assert report["tagged"] is True, "导出文件必须带可导回标记，否则说明导的不是这份台账"

    md2 = MasterData()
    apply_tables(md2, tables)
    assert md2.overview() == md.overview()


def test_roundtrip_keeps_time_range_and_capacity(tmp_path):
    """教务给定的时间范围（星期/节次/周次）+ 容量 是排课输入，往返不能丢。"""
    md = _base_md()
    p = _dump(md, tmp_path / "wb.xlsx")
    md2 = MasterData()
    tables, _ = parse_workbook(p)
    apply_tables(md2, tables)

    c = md2.classes["一班"]
    assert c["weekday"] == "周四" and c["time_slots"] == {"567"}
    assert c["week_min"] == 1 and c["week_max"] == 16
    assert md2.rooms["A410"]["capacity"] == 30


# ---------------------------------------------------------------- 2. 反向验证

def test_edited_cells_actually_land_in_master_data(tmp_path):
    """改过的格子必须真的写回 —— 导入若是 no-op，本用例必须红。"""
    md = _base_md()
    before = md.overview()
    p = _dump(md, tmp_path / "wb.xlsx")

    wb = openpyxl.load_workbook(p)
    # 实验项目页：给「扭摆法」加一位教师
    ws = wb["实验项目"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).strip() == "扭摆法":
            row[2].value = "吴老师、新来的老师"
    # 场地页：改容量
    ws = wb["场地"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).strip() == "A410":
            row[1].value = 60
    # 班级页：加一个组
    ws = wb["班级"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).strip() == "一班":
            row[1].value = "1组、2组、3组"
    wb.save(p)

    tables, _ = parse_workbook(p)
    apply_tables(md, tables)

    assert "新来的老师" in md.experiments["扭摆法"]["teachers"], "教师没写回"
    assert "新来的老师" in md.teachers, "反向引用没建"
    assert md.rooms["A410"]["capacity"] == 60, "容量没写回"
    assert "3组" in md.classes["一班"]["groups"], "学生组没写回"
    assert md.overview() != before, "导入前后完全一样 —— 说明编辑根本没生效（no-op 假绿）"


# ---------------------------------------------------------------- 3. 统计列防线

def test_stat_columns_never_become_entities(tmp_path):
    """老统计台账（「班级数」=5）导回来，不许造出一个名叫「5」的班级。"""
    p = tmp_path / "stat.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["实验项目", "班级数", "学生组数", "课次(格)", "授课教师", "使用场地", "覆盖周次"])
    ws.append(["扭摆法", 5, 10, 30, "吴老师", "A410", "第1-16周（16周）"])
    ws.append(["声速测量", 3, 6, 18, "林老师", "A411", "第1-16周（16周）"])
    wb.save(p)

    md = MasterData()
    tables, report = parse_workbook(str(p))
    apply_tables(md, tables)

    assert "5" not in md.classes, "统计列「班级数」被当成班级名导入了"
    assert "10" not in md.classes
    assert set(md.experiments) == {"扭摆法", "声速测量"}
    assert md.experiments["扭摆法"]["teachers"] == {"吴老师"}
    assert md.experiments["扭摆法"]["rooms"] == {"A410"}
    # 必须**说出口**，不能悄悄忽略
    assert any("统计列" in w for w in report["warnings"]), report["warnings"]
    assert "classes" not in (report["tables"]["experiments"].get("recognized") or [])


# ---------------------------------------------------------------- 4. 认不出就报

def test_unrecognized_headers_are_reported_not_silently_half_imported(tmp_path):
    p = tmp_path / "junk.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["编号", "备注"])
    ws.append([1, "随便写点什么"])
    wb.save(p)

    md = MasterData()
    tables, report = parse_workbook(str(p))
    stats = apply_tables(md, tables)

    assert tables["experiments"] == []
    assert md.experiments == {}
    assert sum(s["rows"] for s in stats.values()) == 0
    assert any("没认出" in w for w in report["warnings"]), report["warnings"]


def test_missing_sheet_is_reported(tmp_path):
    """只给一张「场地」页，其余四张缺页 → 报告里 found=False，不假装导了。"""
    p = tmp_path / "only_rooms.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "场地"
    ws.append(["场地", "容量"])
    ws.append(["A410", 30])
    wb.save(p)

    md = MasterData()
    tables, report = parse_workbook(str(p))
    apply_tables(md, tables)

    assert md.rooms == {"A410": {"capacity": 30}}
    assert report["tables"]["experiments"]["found"] is False
    assert report["tables"]["classes"]["found"] is False


# ---------------------------------------------------------------- 多值分隔

def test_multi_value_separators_all_accepted(tmp_path):
    """教秘手里的表分隔符五花八门：中英文逗号、斜杠、分号都得认。"""
    p = tmp_path / "sep.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["实验项目", "授课教师", "使用场地", "适用班级"])
    ws.append(["扭摆法", "吴老师,林老师", "A410/A411", "一班;二班"])
    wb.save(p)

    md = MasterData()
    tables, _ = parse_workbook(str(p))
    apply_tables(md, tables)
    assert md.experiments["扭摆法"]["teachers"] == {"吴老师", "林老师"}
    assert md.experiments["扭摆法"]["rooms"] == {"A410", "A411"}
    assert md.experiments["扭摆法"]["classes"] == {"一班", "二班"}


def test_blank_cells_do_not_overwrite_existing_values(tmp_path):
    """留空 = 不修改（不是清空）：教秘只填想改的那一列。"""
    md = _base_md()
    p = _dump(md, tmp_path / "wb.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["班级"]
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value).strip() == "一班":
            row[2].value = ""      # 星期留空
            row[4].value = ""      # 起始周留空
    wb.save(p)

    tables, _ = parse_workbook(p)
    apply_tables(md, tables)
    c = md.classes["一班"]
    assert c["weekday"] == "周四", "留空不该清掉已有值"
    assert c["week_min"] == 1


def test_real_schedule_roundtrip_does_not_shatter_class_names(tmp_path):
    """真实 433 条表往返：班级名自带顿号，一个都不能被切碎。

    合成数据测不出这条 —— 真实表里有 5 个班名自带顿号
    （「2024级应用化学1、2班」「24级电子信息工程1、2班」…）。用顿号当多值分隔符
    导出再导回，17 个班会变成 29 个，凭空多出「2班」「3班」「2」这种碎名。
    """
    real = os.path.join(os.path.dirname(__file__), "..", "..", "schedule.xlsx")
    if not os.path.exists(real):
        pytest.skip("仓库根没有 schedule.xlsx")

    db = ScheduleDatabase()
    db.load_from_excel(real)
    md = MasterData().build_from_entries(db.entries)
    before = md.overview()

    p = tmp_path / "real.xlsx"
    build_workbook(md).save(p)
    md2 = MasterData()
    tables, report = parse_workbook(p)
    apply_tables(md2, tables)

    assert len(md2.classes) == len(md.classes) == 17, (
        f"班级数变了（{len(md.classes)} → {len(md2.classes)}）：班名被切碎了")
    assert "2024级应用化学1、2班" in md2.classes
    assert "24级电子信息工程1、2班" in md2.classes
    assert not any(re.fullmatch(r"\d{1,2}班?", c) for c in md2.classes), \
        f"出现了碎班级名：{[c for c in md2.classes if re.fullmatch(r'1,2班?'.replace('1,2', r'\d{1,2}'), c)]}"
    assert md2.overview() == before


def test_top_separator_wins_over_dunhao(tmp_path):
    """一行里既有分号又有顿号 → 只有分号当分隔符（顿号是班名的一部分）。"""
    p = tmp_path / "sep2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["实验项目", "适用班级"])
    ws.append(["扭摆法", "2024级应用化学1、2班;2024级制药工程"])
    wb.save(p)

    md = MasterData()
    apply_tables(md, _parse_only(str(p)))
    assert md.experiments["扭摆法"]["classes"] == {"2024级应用化学1、2班", "2024级制药工程"}


def test_handwritten_dunhao_class_names_are_merged(tmp_path):
    """教秘手写顿号分隔的多班级 → 残片要拼回原班名，不能当新班级。"""
    p = tmp_path / "sep3.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["实验项目", "适用班级"])
    ws.append(["扭摆法", "2024级服装设计与工程1、2、3班、2024级制药工程"])
    wb.save(p)

    md = MasterData()
    apply_tables(md, _parse_only(str(p)))
    assert md.experiments["扭摆法"]["classes"] == {
        "2024级服装设计与工程1、2、3班", "2024级制药工程"}
    # 反过来：合法的「1班、2班」列举不能被合并成一个
    assert _split_multi("1班、2班") == ["1班", "2班"]


# ---------------------------------------------------------------- 端点 + 落盘

@pytest.fixture
def store_path(tmp_path, monkeypatch):
    p = str(tmp_path / "master_store.json")
    monkeypatch.setattr(master_store, "STORE_PATH", p)
    return p


def test_import_endpoint_applies_and_survives_restart(tmp_path, monkeypatch, store_path):
    """走真实 HTTP：导出 → 改 → 导回 → 落盘 → 模拟重启（沙箱回收）→ 数据还在。"""
    from fastapi.testclient import TestClient

    from app.api import reschedule as api
    from app.main import app
    from app.security import issue_admin_token

    db = ScheduleDatabase()
    db.entries = [_entry(1, "周四", "567", "一班", "1组", "A410")]
    db.loaded = True
    db._filepath = str(tmp_path / "schedule.xlsx")
    open(db._filepath, "wb").close()
    monkeypatch.setattr(api, "_db", db)
    monkeypatch.setattr(api, "_master_data", MasterData())
    monkeypatch.setattr(api, "_master_tasks", [])
    monkeypatch.setattr(api, "_schedule_fp", master_store.schedule_fingerprint(db._filepath))
    monkeypatch.setattr(api, "_DATA_DIR", str(tmp_path))
    api._ensure_master_data()

    headers = {"Authorization": f"Bearer {issue_admin_token('admin')['token']}"}
    client = TestClient(app)

    # 导出台账 → 改一格 → 上传
    r = client.get("/api/reschedule/export/master-workbook")
    assert r.status_code == 200, r.text[:200]
    src = tmp_path / "upload.xlsx"
    src.write_bytes(r.content)

    wb = openpyxl.load_workbook(src)
    ws = wb["实验项目"]
    ws.append(["光纤传感", "大学物理实验", "鄢老师", "A502", "一班"])
    wb.save(src)

    r = client.post("/api/reschedule/master/import",
                    files={"file": ("master.xlsx", src.read_bytes(),
                                    "application/vnd.openxmlformats-officedocument."
                                    "spreadsheetml.sheet")},
                    headers=headers)
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert "光纤传感" in {e["name"] for e in body["data"]["experiments"]}, body.get("message")

    # ── 模拟重启：模块态归零 ──
    monkeypatch.setattr(api, "_master_data", MasterData())
    monkeypatch.setattr(api, "_master_tasks", [])
    api._ensure_master_data()
    assert "光纤传感" in api._master_data.experiments, "导入的台账重启后丢了（持久化没接上）"
    assert "鄢老师" in api._master_data.teachers


def test_import_endpoint_rejects_non_excel(tmp_path, monkeypatch, store_path):
    from fastapi.testclient import TestClient

    from app.api import reschedule as api
    from app.main import app
    from app.security import issue_admin_token

    db = ScheduleDatabase()
    db.entries = [_entry(1, "周四", "567", "一班", "1组", "A410")]
    db.loaded = True
    db._filepath = str(tmp_path / "schedule.xlsx")
    open(db._filepath, "wb").close()
    monkeypatch.setattr(api, "_db", db)
    monkeypatch.setattr(api, "_master_data", MasterData())
    monkeypatch.setattr(api, "_schedule_fp", master_store.schedule_fingerprint(db._filepath))

    headers = {"Authorization": f"Bearer {issue_admin_token('admin')['token']}"}
    r = TestClient(app).post(
        "/api/reschedule/master/import",
        files={"file": ("master.txt", b"not an excel", "text/plain")},
        headers=headers)
    assert r.status_code == 400


def test_import_endpoint_requires_admin(tmp_path, monkeypatch, store_path):
    """免鉴权白名单里没有它 —— 主数据是可写实体，匿名不能改。"""
    from fastapi.testclient import TestClient

    from app.api import reschedule as api
    from app.main import app

    db = ScheduleDatabase()
    db.entries = [_entry(1, "周四", "567", "一班", "1组", "A410")]
    db.loaded = True
    monkeypatch.setattr(api, "_db", db)
    monkeypatch.setattr(api, "_master_data", MasterData())

    src = tmp_path / "up.xlsx"
    build_workbook(MasterData()).save(src)
    r = TestClient(app).post(
        "/api/reschedule/master/import",
        files={"file": ("master.xlsx", src.read_bytes(),
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet")})
    assert r.status_code in (401, 403), r.text[:200]
