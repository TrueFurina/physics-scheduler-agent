"""导入向导：预览（只读）→ 确认导入 —— R-63 数据导入向导

要锁的洞
--------
1. **预览绝不写入**：两个预览端点跑完，主数据五张表、`_master_tasks`、
   `db.pending` 必须原样。反向验证（R-33 的用法）：把"预览其实偷偷 apply 了"
   这个 bug 放回代码，本文件相关用例必须红 —— 只断言"接口返回 success"
   会假绿（预览成功≠预览无副作用）。
2. **预览数字 == 导入数字**：同一份文件，预览报的 rows 之和与导入 stats 的
   rows 之和相等。预览与写库共用 `parse_workbook` / `_row_to_rec` /
   `parse_master_tasks`，这条锁的是"两套解析代码慢慢漂移"的回归。
3. **映射必须可见**：字段 ← 列名、没认出的列、统计列忽略、样本行，都要在
   报告里说出口（R-34：认不出就报，不静默半导入）。
4. **鉴权边界与写入一致**：免鉴权白名单里没有预览端点。
5. **清单预览**：列名千奇百怪也能映射；缺「班级/实验项目」→ 400 说清缺什么；
   预览不碰 db.pending。
"""

import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl

from app import master_store
from app.master_data import MasterData
from app.master_workbook import preview_workbook
from app.utils.excel_parser import (PendingTask, ScheduleDatabase, ScheduleEntry,
                                    parse_master_tasks, preview_master_tasks)


def _entry(week, weekday, slot, cls, group, room, teacher="吴老师", experiment="扭摆法"):
    return ScheduleEntry(week, weekday, slot, cls, group, experiment, room, teacher)


def _base_md() -> MasterData:
    md = MasterData()
    md.build_from_entries([
        _entry(1, "周四", "567", "一班", "1组", "A410", teacher="吴老师", experiment="扭摆法"),
        _entry(2, "周四", "567", "一班", "2组", "A411", teacher="林老师", experiment="扭摆法"),
        _entry(3, "周五", "678", "二班", "1组", "A410", teacher="吴老师", experiment="声速测量"),
    ])
    md.upsert_room("A410", capacity=30)
    md.upsert_class("一班", groups=["1组", "2组"], weekday="周四",
                    time_slots=["567"], week_min=1, week_max=16)
    return md


def _dump(md, path):
    from app.master_workbook import build_workbook
    build_workbook(md).save(path)
    return str(path)


# ---------------------------------------------------------------- 单元：台账预览

def test_preview_workbook_maps_columns_and_samples(tmp_path):
    """导出的台账再预览：字段←列名、样本行、行数都要对得上。"""
    md = _base_md()
    p = _dump(md, tmp_path / "wb.xlsx")

    rep = preview_workbook(p)
    assert rep["tagged"] is True
    exp = rep["tables"]["experiments"]
    assert exp["found"] is True and exp["rows"] == 2
    # 列名映射：字段 ← 导出用的表头
    assert exp["columns"]["name"] == "实验项目"
    assert exp["columns"]["teachers"] == "授课教师"
    assert exp["columns"]["classes"] == "适用班级"
    assert exp["unrecognized"] == [], "自己导出的台账不该有认不出的列"
    assert exp["stat_dropped"] == []
    # 样本行与写库同一段代码：名字对得上（overview 的排序不假设）
    assert {s["name"] for s in exp["samples"]} == {"扭摆法", "声速测量"}
    niunai = next(s for s in exp["samples"] if s["name"] == "扭摆法")
    assert niunai["teachers"] == ["吴老师", "林老师"]


def test_preview_reports_unrecognized_and_stat_columns(tmp_path):
    """外校格式：认不出的列、统计列都要在报告里说出口，不当实体名。"""
    p = tmp_path / "foreign.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["编号", "项目名称", "班级数", "任课教师", "备注"])
    ws.append(["E01", "迈克耳孙干涉", 5, "鄢老师", "演示用"])
    ws.append(["E02", "密立根油滴", 3, "陈老师", ""])
    wb.save(p)

    rep = preview_workbook(str(p))
    exp = rep["tables"]["experiments"]
    assert exp["columns"]["name"] == "项目名称", "外校列名要走别名映射"
    assert exp["columns"]["teachers"] == "任课教师"
    assert "编号" in exp["unrecognized"] and "备注" in exp["unrecognized"], \
        "认不出的列必须列出来，否则教秘不知道它们被丢了"
    assert "班级数" not in exp["unrecognized"], \
        "统计列是认出来后主动忽略的，不该又算进未识别列（同一列报两次）"
    assert "班级数" in exp["stat_dropped"], "「班级数」是统计列，不能当班级清单"
    assert any("统计列" in w for w in rep["warnings"]), rep["warnings"]
    # 统计列的值不许以任何形式漏进样本行（5/3 不是班级名）
    assert "classes" not in exp["samples"][0], "统计列被剔除了，样本里不该有班级字段"
    leaked = [v for s in exp["samples"] for v in s.values() if str(v) in ("5", "3")]
    assert not leaked, f"统计列的值漏进样本：{leaked}"


def test_preview_missing_required_column_skips_table_with_reason(tmp_path):
    """表头认不出名称列 → 整页跳过并说清缺什么（R-34：不静默半导入）。"""
    p = tmp_path / "junk.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["编号", "备注"])
    ws.append([1, "随便写点什么"])
    wb.save(p)

    rep = preview_workbook(str(p))
    t = rep["tables"]["experiments"]
    assert t["skipped"] is True and t["rows"] == 0
    assert any("没认出" in w for w in rep["warnings"]), rep["warnings"]


def test_preview_accepts_file_like_object(tmp_path):
    """端点走 BytesIO 解析（预览不落盘）—— 这里锁 file-like 入参可用。"""
    md = _base_md()
    p = _dump(md, tmp_path / "wb.xlsx")
    with open(p, "rb") as f:
        rep = preview_workbook(io.BytesIO(f.read()))
    assert rep["tables"]["experiments"]["rows"] == 2


# ---------------------------------------------------------------- 单元：清单预览

def test_tasks_preview_maps_alias_headers(tmp_path):
    """清单列名千奇百怪（行政班级/项目名称/实验室）也要映射出来。"""
    p = tmp_path / "master.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目清单"
    ws.append(["课程名称", "行政班级", "项目名称", "组别", "指导教师", "实验室", "备注"])
    ws.append(["大学物理实验", "25级新专业", "迈克耳孙干涉", "1组", "吴老师", "A410", "x"])
    ws.append(["大学物理实验", "25级新专业", "声速测量", "2组", "陈老师", "A411", ""])
    wb.save(p)

    rep = preview_master_tasks(str(p))
    assert rep["tasks"] == 2
    assert rep["columns"]["class_name"] == "行政班级"
    assert rep["columns"]["experiment"] == "项目名称"
    assert rep["columns"]["room"] == "实验室"
    assert "备注" in rep["unrecognized_columns"]
    assert rep["samples"][0]["experiment"] == "迈克耳孙干涉"
    assert rep["samples"][0]["teacher"] == "吴老师"
    assert rep["blank_or_incomplete_rows"] == 0


def test_tasks_preview_missing_required_columns_raises(tmp_path):
    """认不出班级/实验项目 → 抛错说清缺什么（端点转 400）。"""
    p = tmp_path / "bad.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["序号", "备注"])
    ws.append([1, "甲"])
    wb.save(p)

    with pytest.raises(ValueError) as ei:
        preview_master_tasks(str(p))
    assert "班级" in str(ei.value) and "实验项目" in str(ei.value)


def test_parse_master_tasks_report_is_structured(tmp_path):
    """写库路径的报告同样带结构化 samples（向导 UI 与 /upload-master 共用）。"""
    p = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目清单"
    ws.append(["课程", "班级", "实验项目", "学生组", "授课教师", "场地"])
    ws.append(["大学物理实验", "25级新专业", "迈克耳孙干涉", "1组", "吴老师", "A410"])
    wb.save(p)

    _tasks, report = parse_master_tasks(str(p))
    assert report["tasks"] == 1
    assert report["samples"] == [{
        "class_name": "25级新专业", "experiment": "迈克耳孙干涉", "group": "1组",
        "teacher": "吴老师", "room": "A410", "course": "大学物理实验",
        "week": None, "weekday": "",
    }]
    assert report["columns"]["class_name"] == "班级"


# ---------------------------------------------------------------- HTTP：只读性 + 鉴权

@pytest.fixture
def api_env(tmp_path, monkeypatch):
    """一套可跑 HTTP 的最小环境（与 test_master_workbook 同模式）。"""
    from app.api import reschedule as api

    db = ScheduleDatabase()
    db.entries = [_entry(1, "周四", "567", "一班", "1组", "A410")]
    db.loaded = True
    db._filepath = str(tmp_path / "schedule.xlsx")
    open(db._filepath, "wb").close()
    monkeypatch.setattr(api, "_db", db)
    monkeypatch.setattr(api, "_master_data", MasterData())
    monkeypatch.setattr(api, "_master_tasks", [])
    monkeypatch.setattr(api, "_schedule_fp", master_store.schedule_fingerprint(db._filepath))
    monkeypatch.setattr(api, "_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(master_store, "STORE_PATH", str(tmp_path / "master_store.json"))
    return api


def _foreign_workbook_bytes() -> bytes:
    """一份外校格式台账：新项目 + 统计列 + 认不出的列。"""
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目"
    ws.append(["编号", "项目名称", "班级数", "任课教师", "备注"])
    ws.append(["E01", "向导预览项目（可删）", 5, "向导预览教师（可删）", "演示用"])
    ws2 = wb.create_sheet("场地")
    ws2.append(["场地", "容量"])
    ws2.append(["A999", 40])
    wb.save(buf)
    return buf.getvalue()


def _admin_headers():
    from app.security import issue_admin_token
    return {"Authorization": f"Bearer {issue_admin_token('admin')['token']}"}


def _post_file(client, path, content: bytes, token=None):
    return client.post(path,
                       files={"file": ("up.xlsx", content,
                                       "application/vnd.openxmlformats-officedocument."
                                       "spreadsheetml.sheet")},
                       headers=token or {})


def test_preview_endpoint_writes_nothing_but_import_does(api_env, tmp_path):
    """核心反向验证：预览跑完主数据必须原样；同一份文件随后能真的导入。

    若预览偷偷 apply 了，「marker 不在 overview 里」这条会红 —— 只断言
    "接口 200"是假绿（预览成功≠无副作用，R-25 铁律）。
    """
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    content = _foreign_workbook_bytes()

    r = _post_file(client, "/api/reschedule/master/import-preview", content, _admin_headers())
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["kind"] == "workbook"
    exp = body["data"]["tables"]["experiments"]
    assert exp["rows"] == 1 and exp["columns"]["name"] == "项目名称"
    assert "向导预览项目（可删）" in [s["name"] for s in exp["samples"]]

    # 预览之后：主数据里不许出现任何来自这份文件的东西
    ov = client.get("/api/reschedule/master/overview").json()["data"]
    names = {e["name"] for e in ov["experiments"]}
    assert "向导预览项目（可删）" not in names, "预览写入了主数据（只读承诺被破坏）"
    assert "A999" not in {r_["name"] for r_ in ov["rooms"]}

    # 同一份文件确认导入：必须真的进得去（证明上面的"没写入"不是文件本身导不进）
    r2 = _post_file(client, "/api/reschedule/master/import", content, _admin_headers())
    assert r2.status_code == 200, r2.text[:300]
    ov2 = client.get("/api/reschedule/master/overview").json()["data"]
    assert "向导预览项目（可删）" in {e["name"] for e in ov2["experiments"]}
    assert "向导预览教师（可删）" in {t["name"] for t in ov2["teachers"]}
    assert r2.json()["stats"]["experiments"]["added"] >= 1


def test_preview_rows_equal_import_rows(api_env):
    """预览报的行数之和 == 导入 stats 的行数之和（同一份解析代码）。"""
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    content = _foreign_workbook_bytes()
    prev = _post_file(client, "/api/reschedule/master/import-preview",
                      content, _admin_headers()).json()["data"]
    prev_rows = sum(t.get("rows", 0) for t in prev["tables"].values())

    imp = _post_file(client, "/api/reschedule/master/import",
                     content, _admin_headers()).json()
    imp_rows = sum(s.get("rows", 0) for s in imp["stats"].values())
    assert prev_rows == imp_rows == 2, f"预览 {prev_rows} 行 vs 导入 {imp_rows} 行"


def test_tasks_preview_writes_nothing_but_upload_does(api_env):
    """清单预览不碰 _master_tasks；确认上传后才进待排任务。"""
    from fastapi.testclient import TestClient

    from app.main import app

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目清单"
    ws.append(["课程", "班级", "实验项目", "学生组"])
    ws.append(["大学物理实验", "25级新专业", "迈克耳孙干涉", "1组"])
    ws.append(["大学物理实验", "25级新专业", "声速测量", "2组"])
    wb.save(buf)
    content = buf.getvalue()

    client = TestClient(app)
    r = _post_file(client, "/api/reschedule/master/tasks-preview", content, _admin_headers())
    assert r.status_code == 200, r.text[:300]
    body = r.json()
    assert body["kind"] == "tasks"
    assert body["data"]["tasks"] == 2
    assert body["data"]["columns"]["class_name"] == "班级"

    st = client.get("/api/reschedule/master-status").json()["data"]
    assert st["master_tasks"] == 0, "预览写入了待排任务（只读承诺被破坏）"

    r2 = _post_file(client, "/api/reschedule/upload-master", content, _admin_headers())
    assert r2.status_code == 200, r2.text[:300]
    st2 = client.get("/api/reschedule/master-status").json()["data"]
    assert st2["master_tasks"] == 2
    assert "25级新专业" in st2["master_classes"]


def test_tasks_preview_leaves_db_pending_untouched(api_env):
    """安排表里已有的待排行，预览清单不能动它。"""
    from fastapi.testclient import TestClient

    from app.main import app

    api_env._db.pending = [PendingTask("一班", "待排实验", "1组")]
    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "实验项目清单"
    ws.append(["课程", "班级", "实验项目", "学生组"])
    ws.append(["大学物理实验", "25级新专业", "迈克耳孙干涉", "1组"])
    wb.save(buf)

    client = TestClient(app)
    r = _post_file(client, "/api/reschedule/master/tasks-preview",
                   buf.getvalue(), _admin_headers())
    assert r.status_code == 200
    st = client.get("/api/reschedule/master-status").json()["data"]
    assert st["pending_in_schedule"] == 1
    assert [t.experiment for t in api_env._db.pending] == ["待排实验"]


def test_tasks_preview_missing_required_columns_400(api_env):
    """缺「班级/实验项目」列 → 400，且消息说清缺什么（与 upload-master 同一句）。"""
    from fastapi.testclient import TestClient

    from app.main import app

    buf = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["序号", "备注"])
    ws.append([1, "甲"])
    wb.save(buf)

    r = _post_file(TestClient(app), "/api/reschedule/master/tasks-preview",
                   buf.getvalue(), _admin_headers())
    assert r.status_code == 400
    assert "班级" in r.json()["detail"] and "实验项目" in r.json()["detail"]


def test_preview_endpoints_require_admin(api_env):
    """免鉴权白名单里没有预览端点（与 /master/import 同一边界）。"""
    from fastapi.testclient import TestClient

    from app.main import app

    content = _foreign_workbook_bytes()
    for path in ("/api/reschedule/master/import-preview",
                 "/api/reschedule/master/tasks-preview"):
        r = _post_file(TestClient(app), path, content)
        assert r.status_code in (401, 403), f"{path} 匿名可达：{r.text[:200]}"


def test_preview_endpoints_reject_non_excel(api_env):
    """非 Excel 内容 → 400（openpyxl 解析失败即拒，与导入端点一致）。"""
    from fastapi.testclient import TestClient

    from app.main import app

    for path in ("/api/reschedule/master/import-preview",
                 "/api/reschedule/master/tasks-preview"):
        r = _post_file(TestClient(app), path, b"not an excel", _admin_headers())
        assert r.status_code == 400, path
