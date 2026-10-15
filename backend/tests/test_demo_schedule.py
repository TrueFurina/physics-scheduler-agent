"""脱敏演示排课表（schedule.demo.xlsx）回归测试。

该文件是可随项目发布的版本，必须同时满足：
  1. 不含任何 PII（真实姓名 / 手机号）
  2. 教学结构与真实表等价（周次/星期/时段/班级/组别/项目/场地）
  3. 解析器读得动，记录数与真实表一致

真实表 schedule.xlsx 被 .gitignore 排除，缺失时跳过等价性比对，但
PII 与可解析性检查仍然执行。
"""

import os
import re
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.utils.excel_parser import ScheduleDatabase

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(__file__)))
DEMO_PATH = os.path.join(REPO_ROOT, "schedule.demo.xlsx")
REAL_PATH = os.path.join(REPO_ROOT, "schedule.xlsx")

PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")

# 真实表实测出现的教师名单（用于确认脱敏后不再出现）
KNOWN_REAL_NAMES = [
    "吴琳", "吴琼", "张妹玉", "林晨", "林珠云", "柯达", "洪锦泉",
    "翁铭华", "肖刚", "郑辉", "鄢仁文", "钟蔚", "陈胜钰",
]

pytestmark = pytest.mark.skipif(not os.path.isfile(DEMO_PATH), reason="缺少 schedule.demo.xlsx")


def test_demo_exists_and_loads():
    db = ScheduleDatabase()
    db.load_from_excel(DEMO_PATH)
    assert len(db.entries) > 0
    assert db.loaded is True


def test_no_phone_numbers_leaked():
    import openpyxl

    wb = openpyxl.load_workbook(DEMO_PATH, data_only=True)
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v and PHONE_RE.search(str(v)):
                    hits.append((ws.title, str(v)[:30]))
    assert hits == [], f"脱敏表仍含手机号：{hits}"


def test_no_real_teacher_names():
    import openpyxl

    wb = openpyxl.load_workbook(DEMO_PATH, data_only=True)
    hits = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if not v:
                    continue
                for name in KNOWN_REAL_NAMES:
                    if name in str(v):
                        hits.add((ws.title, name))
    assert hits == set(), f"脱敏表仍含真实姓名：{sorted(hits)}"


def test_teachers_are_anonymized():
    db = ScheduleDatabase()
    db.load_from_excel(DEMO_PATH)
    teachers = db.get_all_teachers()
    assert teachers, "脱敏表应有教师字段"
    assert all(t.startswith("教师") for t in teachers), f"教师未化名：{teachers}"


def test_sensitive_sheets_removed():
    import openpyxl

    wb = openpyxl.load_workbook(DEMO_PATH)
    for sheet in ("教师联系表", "大学物理实验安排表 (周) ", "大学物理实验安排表 (班级) "):
        assert sheet in wb.sheetnames, f"缺失解析器需要的表：{sheet}"
    for sheet in ("学生所在系系办电话", "任课老师安排表"):
        assert sheet not in wb.sheetnames, f"敏感表未删除：{sheet}"


@pytest.mark.skipif(not os.path.isfile(REAL_PATH), reason="真实表不存在（.gitignore 排除），跳过等价性比对")
def test_structural_equivalence_with_real_schedule():
    """脱敏只应改姓名，教学结构必须与真实表逐一相同。"""
    import collections

    real, demo = ScheduleDatabase(), ScheduleDatabase()
    real.load_from_excel(REAL_PATH)
    demo.load_from_excel(DEMO_PATH)

    assert len(real.entries) == len(demo.entries)

    def sig(db):
        return collections.Counter(
            (e.week, e.weekday, e.time_slot, e.class_name, e.group, e.experiment, e.room)
            for e in db.entries
        )

    assert sig(real) == sig(demo), "脱敏改动了教学结构，需检查替换范围"


# ---------------------------------------------------------------------------
# 默认路径解析契约：
#   SCHEDULE_PATH 显式覆盖  >  真实表（存在时）  >  脱敏表（回退）
#
# 历史说明：早期版本把「默认必须是脱敏表」当作隐私不变量，并显式断言
# 「即便真实表就在旁边也要用脱敏表」。该契约已被产品决策取代——本系统面向
# **教学秘书 / 实验教师**，教师姓名与排课编排是必须真实呈现的业务数据
# （详见 backend/app/api/reschedule.py 的路径解析注释）。
# 真正的隐私防线迁移到了接口层：学生个人身份接口挂管理员门禁，由
# test_api_auth.py 守住；排课表里不被解析的联系方式列则由
# test_no_phone_numbers_leaked 守住脱敏表一侧。
# ---------------------------------------------------------------------------


def _resolver(tmp_path, monkeypatch):
    from app.api import reschedule

    monkeypatch.setattr(reschedule, "_REPO_ROOT", str(tmp_path))
    return reschedule._resolve_default_schedule_path


def test_default_path_prefers_real_schedule_when_present(monkeypatch, tmp_path):
    """真实表存在时必须用真实表——教秘演示要看到真实教师姓名。"""
    monkeypatch.delenv("SCHEDULE_PATH", raising=False)
    (tmp_path / "schedule.xlsx").write_bytes(b"x")
    (tmp_path / "schedule.demo.xlsx").write_bytes(b"x")
    assert _resolver(tmp_path, monkeypatch)() == str(tmp_path / "schedule.xlsx")


def test_default_path_falls_back_to_demo_when_real_absent(monkeypatch, tmp_path):
    """真实表缺失（精简部署）时回退脱敏表，保证服务仍能启动。"""
    monkeypatch.delenv("SCHEDULE_PATH", raising=False)
    (tmp_path / "schedule.demo.xlsx").write_bytes(b"x")
    assert _resolver(tmp_path, monkeypatch)() == str(tmp_path / "schedule.demo.xlsx")


def test_schedule_path_env_override_wins(monkeypatch, tmp_path):
    """显式 SCHEDULE_PATH 优先级最高（相对路径按仓库根解析）。"""
    monkeypatch.setenv("SCHEDULE_PATH", "schedule.xlsx")
    assert _resolver(tmp_path, monkeypatch)() == str(tmp_path / "schedule.xlsx")


def test_schedule_path_env_absolute(monkeypatch, tmp_path):
    monkeypatch.setenv("SCHEDULE_PATH", str(tmp_path / "abs.xlsx"))
    assert _resolver(tmp_path, monkeypatch)() == str(tmp_path / "abs.xlsx")
