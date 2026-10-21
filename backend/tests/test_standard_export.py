"""教务两类标准课表导出：班级课表 / 场地课表

为什么要有这两个接口
--------------------
教务在用系统是**理论课**排课工具，它唯一被认可的优势就是"排完能导出班级课程表、
场地课程表"。本系统此前只能导出**新排出来的并行方案**（/parallel/export），
导不出基于现有安排表的标准课表 —— 教秘排完课还得自己整理一遍才能发到班级群、
贴到实验室门口。

这里锁三件事：
  1. 全表覆盖：导出的行数必须等于安排表条数（一条都不能少，也不能重复）；
  2. 无场地记录不能静默丢弃：它们进「（未填场地）」页（R-35 口径：宁可少推荐，
     不可静默丢弃；静默丢弃会让教秘以为全表都被扫到了）；
  3. 排序：周次 → 星期（含周末）→ 节次，倒序会让教秘在 Excel 里找不到课。
"""

import asyncio
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import app.api.reschedule as api
from app.utils.excel_parser import ScheduleDatabase, ScheduleEntry

WD_ORDER = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
SLOT_ORDER = {"MORNING": 0, "567": 1, "678": 2, "91011": 3}


def _entry(week, weekday, slot, cls, group, room, teacher="林老师", experiment="伏安法",
           blocked=None, time_raw=""):
    return ScheduleEntry(week, weekday, slot, cls, group, experiment, room, teacher,
                         time_raw=time_raw, blocked=blocked)


@pytest.fixture()
def make_db(monkeypatch):
    def _make(entries):
        db = ScheduleDatabase()
        db.entries = entries
        db.loaded = True
        monkeypatch.setattr(api, "_db", db)
        return db

    return _make


def _download(resp) -> io.BytesIO:
    """把 StreamingResponse 读回内存（导出接口返回的是流式响应）。"""
    async def _go():
        return b"".join([chunk async for chunk in resp.body_iterator])

    data = asyncio.run(_go())
    assert data[:2] == b"PK", "导出内容不是 xlsx（前两字节不是 zip 魔数）"
    return io.BytesIO(data)


def _rows(wb):
    """所有 sheet 的数据行（跳过表头）。"""
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    return out


# ---------------------------------------------------------------------------
# 班级课表
# ---------------------------------------------------------------------------

def test_class_schedule_covers_every_entry(make_db):
    """每个班一页，行数之和 == 安排表条数。"""
    make_db([
        _entry(1, "周一", "567", "一班", "1组", "A410"),
        _entry(1, "周一", "567", "一班", "2组", "A411"),
        _entry(2, "周日", "91011", "二班", "1组", "A412-1"),   # 周末也要进
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_class_schedule())))
    rows = _rows(wb)
    assert set(rows) == {"一班", "二班"}
    assert sum(len(v) for v in rows.values()) == 3


def test_class_schedule_rows_are_sorted(make_db):
    """排序必须是 周次 → 星期（含周末）→ 节次。倒序会让教秘在 Excel 里找不到课。"""
    make_db([
        _entry(3, "周五", "91011", "一班", "1组", "A410"),
        _entry(1, "周日", "MORNING", "一班", "1组", "A410"),
        _entry(1, "周六", "567", "一班", "1组", "A410"),
        _entry(1, "周一", "678", "一班", "1组", "A410"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_class_schedule())))
    seq = []
    for r in _rows(wb)["一班"]:
        week = int(str(r[0]).replace("第", "").replace("周", ""))
        seq.append((week, WD_ORDER.index(r[1]), SLOT_ORDER.get(r[2], 99)))
    assert seq == sorted(seq), f"课表没按 周次/星期/节次 排序：{seq}"
    # 反向验证：周末必须排在周五之后（曾经 WEEKDAYS 只有五天，周末会被挤到末尾）
    assert seq[0] == (1, WD_ORDER.index("周一"), SLOT_ORDER["678"])
    assert seq[-1] == (3, WD_ORDER.index("周五"), SLOT_ORDER["91011"])


def test_class_schedule_keeps_cross_slot_original_text(make_db):
    """`14:00-17:00` 要照抄原表写法：写成 13:30-15:45 教秘会以为系统读错了（R-40）。"""
    make_db([
        _entry(12, "周四", "567", "一班", "1组", "A410",
               blocked=["567", "678"], time_raw="14:00-17:00", experiment="设计性实验"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_class_schedule())))
    row = _rows(wb)["一班"][0]
    assert row[3] == "14:00-17:00", f"跨节次记录的上课时间被改写了：{row[3]}"


def test_class_schedule_filter_and_404(make_db):
    make_db([
        _entry(1, "周一", "567", "一班", "1组", "A410"),
        _entry(1, "周二", "567", "二班", "1组", "A411"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_class_schedule("一班"))))
    assert list(_rows(wb)) == ["一班"]

    from fastapi import HTTPException

    with pytest.raises(HTTPException) as ei:
        asyncio.run(api.export_class_schedule("不存在的班"))
    assert ei.value.status_code == 404


# ---------------------------------------------------------------------------
# 场地课表
# ---------------------------------------------------------------------------

ROOMLESS = "（未填场地）"


def test_room_schedule_keeps_roomless_records(make_db):
    """没填场地的记录必须单独成页，不能丢（丢了教秘会以为全表都排好了）。"""
    make_db([
        _entry(1, "周一", "567", "一班", "1组", "A410"),
        _entry(1, "周一", "567", "二班", "1组", "", experiment="绪论课（教室已安排，请查系统）"),
        _entry(1, "周一", "91011", "三班", "1组", "", experiment="绪论课（教室自己申请）"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_room_schedule())))
    rows = _rows(wb)
    assert set(rows) == {"A410", ROOMLESS}
    assert len(rows[ROOMLESS]) == 2, "无场地记录被静默丢弃或合并了"
    assert sum(len(v) for v in rows.values()) == 3


def test_room_schedule_sheet_name_is_excel_safe(make_db):
    """Excel 表名不允许 : \\ / ? * [ ]，且最长 31 字符。"""
    make_db([_entry(1, "周一", "567", "一班", "1组", "A410/2 [临时]")])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_room_schedule())))
    name = wb.sheetnames[0]
    assert len(name) <= 31
    assert not set(name) & set(":\\/?*[]")


def test_room_schedule_404_on_unknown_room(make_db):
    make_db([_entry(1, "周一", "567", "一班", "1组", "A410")])
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as ei:
        asyncio.run(api.export_room_schedule("Z999"))
    assert ei.value.status_code == 404


# ---------------------------------------------------------------------------
# 基础数据台账（教务第一号行动项：梳理实验课程属性 / 重构数据库）
# ---------------------------------------------------------------------------

def test_master_data_has_four_entity_sheets(make_db):
    """四张实体台账：实验项目 / 教师 / 场地 / 班级。"""
    make_db([
        _entry(1, "周一", "567", "一班", "1组", "A410", teacher="林老师", experiment="扭摆法"),
        _entry(2, "周六", "678", "二班", "1组", "A411", teacher="吴老师", experiment="扭摆法"),
        _entry(2, "周日", "MORNING", "二班", "2组", "A411", teacher="吴老师", experiment="光的偏振"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_master_data())))
    assert wb.sheetnames == ["实验项目", "教师", "场地", "班级"]
    proj = _rows(wb)["实验项目"]
    # 同一项目在两个班开 ⇒ 班级数 2，授课教师两人
    twist = [r for r in proj if r[0] == "扭摆法"][0]
    assert twist[1] == 2
    assert set(twist[4].split("、")) == {"林老师", "吴老师"}
    assert len(_rows(wb)["教师"]) == 2
    assert len(_rows(wb)["场地"]) == 2
    assert len(_rows(wb)["班级"]) == 2


def test_master_data_keeps_unnamed_entities(make_db):
    """缺教师/缺场地的记录要归入"未指派/未填"，不能整条丢掉（台账是用来补全的）。"""
    make_db([
        _entry(1, "周一", "567", "一班", "1组", "", teacher="", experiment="绪论课"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_master_data())))
    rows = _rows(wb)
    assert [r[0] for r in rows["教师"]] == ["未指派教师"]
    assert [r[0] for r in rows["场地"]] == []          # 无场地不进场地页（另在场地课表里单列）
    assert [r[0] for r in rows["实验项目"]] == ["绪论课"]
    assert rows["实验项目"][0][5] == "（未填场地）"


def test_master_data_class_weekday_mix_includes_weekend(make_db):
    """星期分布必须含周末且按周一→周日排序（曾经 WEEKDAYS 只有五天，周末被挤到末尾）。"""
    make_db([
        _entry(1, "周日", "567", "一班", "1组", "A410"),
        _entry(2, "周六", "567", "一班", "1组", "A410"),
        _entry(2, "周一", "567", "一班", "1组", "A410"),
    ])
    import openpyxl

    wb = openpyxl.load_workbook(_download(asyncio.run(api.export_master_data())))
    mix = _rows(wb)["班级"][0][5]
    assert mix == "周一 1、周六 1、周日 1", f"星期分布不含周末或顺序不对：{mix}"
